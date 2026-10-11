"""Read any uploaded document into a LoadedDocument with Sections. Owner: Terminal C.

Generic on purpose: nothing here knows about a particular client's documents.

heading_path holds the document's own headings only (Markdown/HTML/DOCX/PDF headings,
"Page N" for PDFs without detectable headings, the sheet name for spreadsheets). The title
is not prepended when the document has no heading of its own, so chunking should add the
title in front when heading_path[0] != title.

Title rule: the first heading, else the file's metadata title, else the file name. A document
that opens with a top-level heading that has no text under it, directly followed by another
heading of the same level (e.g. "# HR Record" then "# Alex Chen"), is treated as label + name:
the title becomes "HR Record: Alex Chen".

OCR (RapidOCR, local and free) reads images and the PDF pages that have no text layer; pages
with a normal text layer are never OCR'd. Run `python -m answer_engine.loaders --prepare-ocr`
during setup / the Hugging Face build so the first upload never waits for the OCR models.
"""
import csv
import hashlib
import io
import re
from collections import Counter
from datetime import date, datetime, time
from functools import lru_cache
from pathlib import Path

from answer_engine.schema import LoadedDocument, Section

SUPPORTED_EXTENSIONS: set[str] = {
    ".md", ".markdown", ".txt", ".pdf", ".docx", ".xlsx", ".csv", ".html", ".htm", ".png", ".jpg", ".jpeg",
}
TABLE_HEADER_EVERY = 30  # spreadsheet rows between repeated header rows
OCR_DPI = 200  # resolution scanned PDF pages are rendered at before OCR
OCR_MIN_SCORE = 0.5  # drop OCR guesses below this confidence
OCR_MAX_STAMP_CHARS = 20  # a page with an image and less text than this (e.g. a page-number stamp) is OCR'd


class NoTextFoundError(ValueError):
    """The file has no readable text, not even after OCR (blank scan, photo without words)."""


def load_file(path: Path, doc_id: str) -> LoadedDocument:
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type '{ext or path.name}'. Supported types: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")
    raw = path.read_bytes()
    file_type = {".markdown": "md", ".htm": "html", ".jpeg": "jpg"}.get(ext, ext[1:])
    title, sections, metadata = _LOADERS[file_type](raw, path)
    return LoadedDocument(
        doc_id=doc_id,
        title=title or path.stem,
        source=doc_id,
        file_type=file_type,
        sections=sections,
        content_hash=hashlib.sha256(raw).hexdigest(),
        metadata={"file_name": path.name, **metadata},
    )


# ---------- shared helpers ----------


class _SectionBuilder:
    """Collects text under the current heading path and emits a Section at each heading or table."""

    def __init__(self):
        self.stack: list[tuple[int, str]] = []
        self.lines: list[str] = []
        self.sections: list[Section] = []
        self.title: str | None = None
        self._headings_seen = 0
        self._content_seen = False

    def heading(self, level: int, text: str):
        text = _clean_inline(text)
        if not text:
            return
        self._flush()
        is_label = (
            self._headings_seen == 1
            and not self._content_seen
            and len(self.stack) == 1
            and self.stack[0][0] == level
        )
        self._headings_seen += 1
        if is_label:
            self.stack[0] = (level, f"{self.stack[0][1]}: {text}")
            self.title = self.stack[0][1]
            return
        while self.stack and self.stack[-1][0] >= level:
            self.stack.pop()
        self.stack.append((level, text))
        if self.title is None:
            self.title = text

    def text(self, line: str):
        if line.strip():
            self._content_seen = True
        self.lines.append(line.rstrip())

    def table(self, markdown: str):
        self._flush()
        if markdown.strip():
            self._content_seen = True
            self.sections.append(Section(heading_path=self._path(), text=markdown, kind="table"))

    def finish(self) -> list[Section]:
        self._flush()
        return self.sections

    def _path(self) -> list[str]:
        return [text for _, text in self.stack]

    def _flush(self):
        text = _tidy("\n".join(self.lines))
        self.lines = []
        if text:
            self.sections.append(Section(heading_path=self._path(), text=text))


def _clean_inline(text: str) -> str:
    text = re.sub(r"\s+", " ", text).strip()
    return text.strip("*_` ").strip()


def _tidy(text: str) -> str:
    text = "\n".join(line.rstrip() for line in text.splitlines())
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    elif isinstance(value, datetime):
        value = value.date().isoformat() if value.time() == time() else value.isoformat(sep=" ")
    elif isinstance(value, (date, time)):
        value = value.isoformat()
    return re.sub(r"\s+", " ", str(value)).strip().replace("|", "\\|")


def _markdown_table(rows: list[list], repeat_header_every: int | None = None) -> str:
    """First row is the header. Optionally repeat it every N body rows so each part keeps its column names."""
    rows = [[_cell(v) for v in row] for row in rows]
    rows = [row for row in rows if any(row)]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    while width > 1 and not any(len(row) >= width and row[width - 1] for row in rows):
        width -= 1
    rows = [(row + [""] * width)[:width] for row in rows]
    header, body = rows[0], rows[1:]
    header_lines = ["| " + " | ".join(header) + " |", "|" + "---|" * width]
    step = repeat_header_every or max(len(body), 1)
    blocks = []
    for start in range(0, max(len(body), 1), step):
        part = body[start : start + step]
        blocks.append("\n".join(header_lines + ["| " + " | ".join(row) + " |" for row in part]))
    return "\n\n".join(blocks)


# ---------- Markdown and plain text ----------

_ATX = re.compile(r"^ {0,3}(#{1,6})(?:\s+(.*?))?\s*#*\s*$")
_FENCE = re.compile(r"^ {0,3}(`{3,}|~{3,})")
_TABLE_SEPARATOR = re.compile(r"^\s*\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)*\|?\s*$")


def _load_markdown(raw: bytes, path: Path):
    lines = _decode(raw).splitlines()
    builder = _SectionBuilder()
    fence = None
    i = 0
    while i < len(lines):
        line = lines[i]
        fence_match = _FENCE.match(line)
        if fence:
            builder.text(line)
            if fence_match and fence_match.group(1)[0] == fence[0] and len(fence_match.group(1)) >= len(fence):
                fence = None
            i += 1
            continue
        if fence_match:
            fence = fence_match.group(1)
            builder.text(line)
            i += 1
            continue
        heading = _ATX.match(line)
        if heading:
            builder.heading(len(heading.group(1)), heading.group(2) or "")
            i += 1
            continue
        # setext heading: a paragraph line underlined with === or ---
        next_line = lines[i + 1] if i + 1 < len(lines) else ""
        if line.strip() and not line.lstrip().startswith(("|", "-", "*", "+", ">")) and re.fullmatch(r" {0,3}(=+|-+)\s*", next_line):
            if not (builder.lines and builder.lines[-1].strip()):
                builder.heading(1 if next_line.strip()[0] == "=" else 2, line)
                i += 2
                continue
        if line.lstrip().startswith("|") and _TABLE_SEPARATOR.match(next_line) and "-" in next_line:
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            builder.table("\n".join(block))
            continue
        builder.text(line)
        i += 1
    return builder.title, builder.finish(), {}


def _load_text(raw: bytes, path: Path):
    text = _tidy(_decode(raw))
    return None, [Section(heading_path=[], text=text)] if text else [], {}


# ---------- HTML ----------

_SKIP_TAGS = {"script", "style", "noscript", "template", "nav", "svg", "head"}
_BLOCK_TEXT_TAGS = {"p", "pre", "blockquote", "dt", "dd", "figcaption", "caption", "address"}


def _load_html(raw: bytes, path: Path):
    from bs4 import BeautifulSoup, NavigableString, Tag

    soup = BeautifulSoup(raw, "html.parser")
    meta_title = soup.title.get_text(" ", strip=True) if soup.title else None
    builder = _SectionBuilder()

    def walk(node):
        for child in node.children:
            if isinstance(child, NavigableString):
                if type(child) is NavigableString and child.strip():
                    builder.text(child.strip())
                continue
            if not isinstance(child, Tag) or child.name in _SKIP_TAGS:
                continue
            name = child.name
            if re.fullmatch(r"h[1-6]", name):
                builder.heading(int(name[1]), child.get_text())
            elif name == "table":
                builder.table(_html_table(child))
            elif name == "li":
                inline = "".join(
                    c.get_text() if isinstance(c, Tag) else str(c)
                    for c in child.children
                    if not (isinstance(c, Tag) and c.name in ("ul", "ol"))
                )
                builder.text("- " + _clean_inline(inline))
                for nested in child.find_all(["ul", "ol"], recursive=False):
                    walk(nested)
            elif name == "br":
                continue
            elif name in _BLOCK_TEXT_TAGS:
                builder.text(child.get_text() if name == "pre" else _clean_inline(child.get_text()))
                builder.text("")
            else:
                walk(child)

    walk(soup.body or soup)
    return builder.title or meta_title, builder.finish(), {}


def _html_table(table) -> str:
    rows = []
    for tr in table.find_all("tr"):
        if tr.find_parent("table") is not table:
            continue  # rows of a nested table
        rows.append([cell.get_text() for cell in tr.find_all(["th", "td"], recursive=False)])
    return _markdown_table(rows)


# ---------- Word ----------


def _load_docx(raw: bytes, path: Path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(io.BytesIO(raw))
    builder = _SectionBuilder()
    for element in document.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "tbl":
            table = Table(element, document)
            builder.table(_markdown_table([[cell.text for cell in _unique_cells(row)] for row in table.rows]))
        elif tag == "p":
            paragraph = Paragraph(element, document)
            style = (paragraph.style.name if paragraph.style is not None else "") or ""
            level = _docx_heading_level(style)
            if level is not None:
                builder.heading(level, paragraph.text)
            elif "list" in style.lower() and paragraph.text.strip():
                builder.text("- " + paragraph.text.strip())
            else:
                builder.text(paragraph.text)
                builder.text("")
    meta_title = (document.core_properties.title or "").strip() or None
    return builder.title or meta_title, builder.finish(), {}


def _docx_heading_level(style: str) -> int | None:
    name = style.lower()
    if name == "title":
        return 0
    match = re.fullmatch(r"heading\s*(\d)", name)
    return int(match.group(1)) if match else None


def _unique_cells(row):
    """python-docx repeats a merged cell once per grid column; keep each cell once."""
    seen, cells = set(), []
    for cell in row.cells:
        if id(cell._tc) not in seen:
            seen.add(id(cell._tc))
            cells.append(cell)
    return cells


# ---------- PDF ----------

_BOLD_FONT = re.compile(r"bold|black|heavy|semibold|demi", re.IGNORECASE)


def _load_pdf(raw: bytes, path: Path):
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError(f"{path.name} is password-protected and cannot be read")

    pages = []  # (page text, {normalised line: (size, bold)})
    ocr_pages = []
    rendered = None  # pypdfium2 document, opened only if some page needs OCR
    for number, page in enumerate(reader.pages, start=1):
        spans: dict[float, list] = {}

        def visit(text, cm, tm, font_dict, font_size, spans=spans):
            if not text.strip():
                return
            scale = abs(tm[3] * cm[3]) or abs(tm[0] * cm[0]) or 1
            y = round(tm[5] * cm[3] + cm[5], 1)
            font = str((font_dict or {}).get("/BaseFont", ""))
            spans.setdefault(y, []).append((text, font_size * scale, bool(_BOLD_FONT.search(font))))

        text = page.extract_text(visitor_text=visit) or ""
        visible = len(re.sub(r"\s+", "", text))
        if visible == 0 or (visible < OCR_MAX_STAMP_CHARS and _page_has_images(page)):
            if rendered is None:
                import pypdfium2

                rendered = pypdfium2.PdfDocument(raw)
            image = rendered[number - 1].render(scale=OCR_DPI / 72).to_pil()
            pages.append((_ocr_image(image), {}))  # no font info, so OCR text is never a heading
            ocr_pages.append(number)
            continue
        styled = {}
        for parts in spans.values():
            key = _normalise("".join(t for t, _, _ in parts))
            if key:
                styled[key] = (max(size for _, size, _ in parts), all(bold for t, _, bold in parts if t.strip()))
        pages.append((text, styled))

    if rendered is not None:
        rendered.close()
    if not any(text.strip() for text, _ in pages):
        raise NoTextFoundError(f"No readable text found in {path.name}, not even with OCR (blank or unreadable scan).")

    heading_levels, body_size = _pdf_heading_levels(pages)
    builder = _SectionBuilder()
    for number, (text, styled) in enumerate(pages, start=1):
        if not heading_levels:
            builder.heading(1, f"Page {number}")
        for line in text.splitlines():
            style = styled.get(_normalise(line))
            level = _pdf_line_level(line, style, heading_levels, body_size)
            if level is not None:
                builder.heading(level, line)
            else:
                builder.text(line)
    meta_title = None
    if reader.metadata and reader.metadata.title:
        meta_title = str(reader.metadata.title).strip() or None
    title = builder.title if heading_levels else None
    return title or meta_title, builder.finish(), {"pages": len(pages), "ocr_pages": ocr_pages}


def _page_has_images(page) -> bool:
    try:
        xobjects = page["/Resources"]["/XObject"]
        return any(xobjects[name].get("/Subtype") == "/Image" for name in xobjects)
    except Exception:
        return False


def _normalise(text: str) -> str:
    return re.sub(r"\s+", "", text).lower()


def _pdf_heading_levels(pages) -> tuple[dict, float]:
    """Map heading font sizes to levels. Bigger than body text = heading; bold at body size = lowest level."""
    sizes = Counter()
    for text, styled in pages:
        for key, (size, _) in styled.items():
            sizes[round(size, 1)] += len(key)
    if not sizes:
        return {}, 0.0
    body = sizes.most_common(1)[0][0]
    candidates = set()
    for text, styled in pages:
        for line in text.splitlines():
            style = styled.get(_normalise(line))
            if style and _looks_like_heading(line) and (style[0] >= body * 1.15 or (style[1] and style[0] >= body * 0.95)):
                candidates.add(round(style[0], 1) if style[0] >= body * 1.15 else "bold")
    big = sorted((s for s in candidates if s != "bold"), reverse=True)
    levels = {size: i + 1 for i, size in enumerate(big)}
    if "bold" in candidates:
        levels["bold"] = len(big) + 1
    return levels, body


def _pdf_line_level(line, style, levels, body) -> int | None:
    if not levels or not style or not _looks_like_heading(line):
        return None
    size, bold = style
    if size >= body * 1.15:
        return levels.get(round(size, 1))
    if bold and size >= body * 0.95:
        return levels.get("bold")
    return None


def _looks_like_heading(line: str) -> bool:
    line = line.strip()
    return 0 < len(line) <= 120 and len(line.split()) <= 15 and not line.endswith((".", ",", ";")) and any(c.isalpha() for c in line)


# ---------- images and OCR ----------


def prepare_ocr() -> None:
    """Download any missing OCR model and load it, so the first scanned upload doesn't wait.
    Call during setup or the Hugging Face build: python -m answer_engine.loaders --prepare-ocr"""
    _ocr_engine()


@lru_cache(maxsize=1)
def _ocr_engine():
    from rapidocr import RapidOCR

    # creating the engine downloads missing models and verifies the bundled ones by checksum
    return RapidOCR(params={"Global.log_level": "error"})


def _load_image(raw: bytes, path: Path):
    from PIL import Image, ImageOps

    image = ImageOps.exif_transpose(Image.open(io.BytesIO(raw))).convert("RGB")
    text = _ocr_image(image)
    if not text:
        raise NoTextFoundError(f"No readable text found in {path.name}.")
    return None, [Section(heading_path=[], text=text)], {"ocr": True}


def _ocr_image(image) -> str:
    """OCR a PIL image; returns the text as lines in reading order (top to bottom, left to right)."""
    import numpy as np

    result = _ocr_engine()(np.asarray(image.convert("RGB")))
    if result.txts is None:
        return ""
    words = []
    for box, text, score in zip(result.boxes, result.txts, result.scores):
        if score >= OCR_MIN_SCORE and text.strip():
            ys, xs = [p[1] for p in box], [p[0] for p in box]
            words.append((min(ys), max(ys), min(xs), text.strip()))
    lines: list[list] = []
    for top, bottom, left, text in sorted(words):
        middle = (top + bottom) / 2
        if lines and lines[-1][0] <= middle <= lines[-1][1]:  # vertically inside the current line
            lines[-1][2].append((left, text))
        else:
            lines.append([top, bottom, [(left, text)]])
    return "\n".join(" ".join(text for _, text in sorted(parts)) for _, _, parts in lines)


# ---------- spreadsheets ----------


def _load_xlsx(raw: bytes, path: Path):
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        sections = []
        for sheet in workbook.worksheets:
            table = _markdown_table([list(row) for row in sheet.iter_rows(values_only=True)], TABLE_HEADER_EVERY)
            if table:
                sections.append(Section(heading_path=[sheet.title], text=table, kind="table"))
        meta_title = (workbook.properties.title or "").strip() or None
        return meta_title, sections, {"sheets": len(workbook.worksheets)}
    finally:
        workbook.close()


def _load_csv(raw: bytes, path: Path):
    text = _decode(raw)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    table = _markdown_table(rows, TABLE_HEADER_EVERY)
    return None, [Section(heading_path=[], text=table, kind="table")] if table else [], {}


_LOADERS = {
    "md": _load_markdown,
    "txt": _load_text,
    "html": _load_html,
    "docx": _load_docx,
    "pdf": _load_pdf,
    "xlsx": _load_xlsx,
    "csv": _load_csv,
    "png": _load_image,
    "jpg": _load_image,
}


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Document loader utilities.")
    parser.add_argument("--prepare-ocr", action="store_true", help="download (if missing) and load the OCR models")
    if parser.parse_args().prepare_ocr:
        prepare_ocr()
        print("OCR models ready.")
    else:
        parser.print_help()
