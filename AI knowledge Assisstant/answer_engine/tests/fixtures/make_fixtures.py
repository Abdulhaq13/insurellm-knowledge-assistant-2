"""Builds the small test documents in this folder. Re-run after changing them:
python answer_engine/tests/fixtures/make_fixtures.py"""
import csv
from datetime import date
from pathlib import Path

import docx
import openpyxl
from fpdf import FPDF

HERE = Path(__file__).parent

MARKDOWN = """# Policy Handbook

# Travel Policy

Intro text before any subheading.

## Booking

Book flights through the portal.

```bash
# this is a comment, not a heading
echo hi
```

### Approval

Managers approve trips over $500.

| Grade | Daily limit |
|---|---|
| Staff | $100 |
| Manager | $150 |

Text after the table.

Expenses
--------

Submit receipts within 30 days.
"""

TEXT = "Meeting notes\n\nWe agreed to ship on Friday.\n\n\n\nAction: write tests.\n"

HTML = """<!doctype html>
<html><head><title>Meta Title</title><style>body {color: red}</style></head>
<body>
<nav>Home | About</nav>
<h1>Product Guide</h1>
<p>Welcome to the <b>guide</b>.</p>
<h2>Plans</h2>
<ul><li>Basic</li><li>Pro<ul><li>Pro Plus</li></ul></li></ul>
<table>
  <tr><th>Plan</th><th>Price</th></tr>
  <tr><td>Basic</td><td>$10</td></tr>
  <tr><td>Pro</td><td>$20 | month</td></tr>
</table>
<h3>Discounts</h3>
<p>Annual billing saves 10%.</p>
<script>var hidden = "do not index";</script>
</body></html>
"""


def make_docx(path: Path):
    document = docx.Document()
    document.core_properties.title = "Core Title"
    document.add_paragraph("Employee Guide", style="Title")
    document.add_paragraph("Opening paragraph.")
    document.add_heading("Leave", level=1)
    document.add_paragraph("Staff get 25 days of leave.")
    document.add_paragraph("Carry over up to 5 days", style="List Bullet")
    document.add_heading("Sick Leave", level=2)
    document.add_paragraph("Tell your manager by 9am.")
    table = document.add_table(rows=3, cols=2)
    for r, row in enumerate([["Years", "Days"], ["0-2", "25"], ["3+", "28"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    document.add_heading("Benefits", level=1)
    document.add_paragraph("Health cover for all staff.")
    document.save(path)


def make_xlsx(path: Path):
    workbook = openpyxl.Workbook()
    sales = workbook.active
    sales.title = "Sales"
    sales.append(["Region", "Month", "Revenue"])
    for i in range(70):
        sales.append([f"R{i % 4}", date(2025, i % 12 + 1, 1), 1000.0 + i])
    staff = workbook.create_sheet("Staff")
    staff.append(["Name", "Role", None])
    staff.append(["Ana", "Engineer", None])
    workbook.create_sheet("Empty")
    workbook.properties.title = "Quarterly Report"
    workbook.save(path)


def make_csv(path: Path):
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["id", "name", "amount"])
        for i in range(65):
            writer.writerow([i, f"item {i}", i * 2.5])


def make_semicolon_csv(path: Path):
    path.write_text("city;population\nOslo;700000\nBergen;290000\n", encoding="utf-8")


def make_pdf_with_headings(path: Path):
    pdf = FPDF()
    pdf.set_title("PDF Meta Title")
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 22)
    pdf.cell(0, 12, "Annual Report", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Revenue", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 11)
    pdf.multi_cell(0, 6, "Revenue grew by 12 percent this year. The growth came from new customers in three regions.")
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(0, 8, "Regional Detail", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 11)
    pdf.multi_cell(0, 6, "The north region contributed most of the new revenue.")
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Costs", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 11)
    pdf.multi_cell(0, 6, "Costs fell by 3 percent because hosting became cheaper.")
    pdf.output(str(path))


def make_plain_pdf(path: Path):
    pdf = FPDF()
    pdf.set_font("Helvetica", "", 11)
    for text in ["First page body text that has no headings at all.", "Second page body text, also plain."]:
        pdf.add_page()
        pdf.multi_cell(0, 6, text)
    pdf.output(str(path))


def make_scanned_pdf(path: Path):
    pdf = FPDF()
    pdf.add_page()
    pdf.rect(20, 20, 100, 60, style="F")  # an image-like page with no text layer
    pdf.output(str(path))


def main():
    (HERE / "sample.md").write_text(MARKDOWN, encoding="utf-8")
    (HERE / "sample.txt").write_text(TEXT, encoding="utf-8")
    (HERE / "sample.html").write_text(HTML, encoding="utf-8")
    make_docx(HERE / "sample.docx")
    make_xlsx(HERE / "sample.xlsx")
    make_csv(HERE / "sample.csv")
    make_semicolon_csv(HERE / "semicolon.csv")
    make_pdf_with_headings(HERE / "headings.pdf")
    make_plain_pdf(HERE / "plain.pdf")
    make_scanned_pdf(HERE / "scanned.pdf")
    print(f"Fixtures written to {HERE}")


if __name__ == "__main__":
    main()
